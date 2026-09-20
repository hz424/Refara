#!/usr/bin/env python3
"""Write a compact reviewer workbook containing every Figure 5 plotted value."""
from __future__ import annotations
import csv
from datetime import datetime
from io import BytesIO
from pathlib import Path
import zipfile
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE=Path(__file__).resolve().parent


def read(name):
    with (HERE/'source_data'/name).open(newline='') as stream:
        return list(csv.DictReader(stream,delimiter='\t'))


def main():
    book=Workbook()
    guide=book.active
    guide.title='Read me'
    guide.append(['Sheet','Contents'])
    guide.append(['5a Global null','All 40 archived global-null conditions; 5,000 Monte Carlo replicates each.'])
    guide.append(['5b Hierarchical null','All five procedures at eight root counts; 2,500 replicates per root count, 499 bootstrap draws.'])
    guide.append(['5c Recovery','All 24 archived recovery summaries, pooled over 10,000 positive and negative simulations.'])
    guide.append(['5c Analytic region','R=1–7: zero recovery follows exactly from sign-test discreteness and 12-direction Holm correction. Blank intervals are intentional.'])
    guide.append(['Intervals','Pointwise two-sided 95% Monte Carlo intervals; Clopper–Pearson for event rates, Hoeffding for mean true-edge recovery.'])
    guide.append(['Units','Rates are proportions (0–1). R counts independent acquisition units; tasks and cells are descendants.'])
    guide.append(['Provenance','R=2,3,4,6 in panel b are the extension. All original R=8,12,20,40 values are preserved. TSV files retain detailed source identifiers.'])
    configurations=[
        ('5a Global null','F2B_GLOBAL_NULL_CALIBRATION_V1.tsv', ['axis_variant_label','root_count_R','events','trials','point_estimate','interval_lower','interval_upper','interval_method']),
        ('5b Hierarchical null','FIGURE_5B_ALL_ROOT_COUNTS.tsv',['root_count_R','procedure_label','events','trials','estimate','interval_lower','interval_upper','invalid_descendant_replicate']),
        ('5c Recovery','F2A_BASE_RECOVERY_POOLED_V1.tsv',['endpoint_label','absolute_standardized_effect','root_count_R','pooled_numerator_or_score_sum','pooled_denominator','point_estimate','interval_lower','interval_upper','interval_method']),
        ('5c Analytic region','FIGURE_5C_ANALYTIC_LOWER_ROOT_REGION.tsv',['independent_roots_R','minimum_one_sided_sign_p','first_holm_threshold','rejection_possible','exact_family_recovery','true_edge_recovery','basis','monte_carlo_replicates','interval_lower','interval_upper'])]
    numeric={'root_count_R','events','trials','point_estimate','estimate','interval_lower','interval_upper','absolute_standardized_effect','pooled_numerator_or_score_sum','pooled_denominator','independent_roots_R','exact_family_recovery','true_edge_recovery'}
    for title,file,columns in configurations:
        sheet=book.create_sheet(title)
        sheet.append(columns)
        for row in read(file):
            values=[]
            for column in columns:
                value=row[column]
                if value=='':value=None
                elif column in numeric:value=float(value)
                values.append(value)
            sheet.append(values)
    for sheet in book:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='3D7180')
        for col in sheet.columns:
            key=col[0].column_letter
            sheet.column_dimensions[key].width=min(90,max(14,max(len(str(c.value or '')) for c in col)+2))
            for cell in col[1:]:
                if isinstance(cell.value,float) and cell.value%1:
                    cell.number_format='0.0000'
    book.properties.creator=''
    book.properties.lastModifiedBy=''
    book.properties.created=datetime(2026,9,11)
    book.properties.modified=datetime(2026,9,11)
    raw=BytesIO()
    book.save(raw)
    # Stable ZIP metadata makes this source-data artifact exactly replayable.
    with zipfile.ZipFile(BytesIO(raw.getvalue())) as source:
        with zipfile.ZipFile(HERE/'source_data/Figure_5_Source_Data.xlsx','w',zipfile.ZIP_DEFLATED,compresslevel=9) as target:
            for name in sorted(source.namelist()):
                info=zipfile.ZipInfo(name,date_time=(2026,9,11,0,0,0))
                info.compress_type=zipfile.ZIP_DEFLATED
                content=source.read(name)
                if name=='docProps/core.xml':
                    content=re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',
                                   rb'\g<1>2026-09-11T00:00:00Z\g<2>',content)
                target.writestr(info,content)


if __name__=='__main__':
    main()
