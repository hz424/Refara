#!/usr/bin/env python3
"""Prepare revised Extended Data Figure 8 source tables and workbook."""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import itertools
import json
from pathlib import Path
import re
import shutil
import zipfile

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

DEPTHS = [8, 16, 32, 64, 128, 135]
PATTERNS = ['S', 'M', 'P', 'O', 'D']
STATE_MODELS = ['additive', 'compositional_ridge', 'CPA_0.8.8']


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_tsv(path):
    return pd.read_csv(path, sep='\t', float_precision='round_trip')


def stable_xlsx(book, destination):
    raw = destination.with_suffix('.temporary.xlsx')
    book.save(raw)
    with zipfile.ZipFile(raw) as archive, zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as output:
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            content = archive.read(name)
            if name == 'docProps/core.xml':
                content = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',
                                 rb'\g<1>2000-01-01T00:00:00Z\g<2>', content)
            output.writestr(info, content)
    raw.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--original-source-dir', required=True, type=Path)
    parser.add_argument('--results-dir', required=True, type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    original_record = json.loads((args.original_source_dir/'SOURCE_DATA_RECEIPT.json').read_text())
    score_record = json.loads((args.results_dir/'score_receipt.json').read_text())
    assert original_record['status'] == 'PASS'
    assert score_record['status'] == 'PASS_DIRECT_EFFECT_COMPARISONS'
    for record, directory in [(original_record, args.original_source_dir), (score_record, args.results_dir)]:
        for name, digest in record['output_sha256'].items():
            assert sha256(directory/name) == digest, name
    source_names = [name for name in original_record['output_sha256'] if name.endswith('.tsv')]
    assert len(source_names) == 9
    for name in source_names:
        shutil.copyfile(args.original_source_dir/name, args.output_dir/name)
    arrays = np.load(args.results_dir/'effect_utilities.npz', allow_pickle=False)
    assert arrays['depths'].tolist() == DEPTHS and arrays['patterns'].tolist() == PATTERNS
    assert arrays['state_models'].tolist() == STATE_MODELS
    utility = arrays['effect_utility']
    assert utility.shape == (55, 30, 6, 5) and np.isfinite(utility).all()
    task_mean = utility.mean(axis=0)
    a_rows = [(allocation, DEPTHS[di], 8*DEPTHS[di], 'direct_effect_ridge', PATTERNS[pi], -task_mean[allocation, di, pi])
              for allocation, di, pi in itertools.product(range(30), range(6), range(5))]
    a_added = pd.DataFrame(a_rows, columns=['allocation', 'depth_per_gemgroup', 'cells_per_reference', 'model', 'pattern', 'scaled_squared_error'])
    a_added_summary = a_added.groupby(['depth_per_gemgroup', 'cells_per_reference', 'model', 'pattern'], sort=False).scaled_squared_error.agg(
        mean='mean', allocation_min='min', allocation_max='max').reset_index()
    for name, frame in [('a_allocations', a_added), ('a_summary', a_added_summary)]:
        # Append bytes so every previous source row is retained verbatim.
        original = (args.output_dir/f'{name}.tsv').read_bytes()
        with (args.output_dir/f'{name}.tsv').open('ab') as stream:
            stream.write(frame.to_csv(sep='\t', index=False, header=False, float_format='%.17g').encode())
        assert (args.output_dir/f'{name}.tsv').read_bytes().startswith(original)
        combined = read_tsv(args.output_dir/f'{name}.tsv')
        retained = combined.loc[combined.model != 'direct_effect_ridge'].reset_index(drop=True)
        pd.testing.assert_frame_equal(retained, read_tsv(args.original_source_dir/f'{name}.tsv'), check_exact=True)
    for result_name, source_name in [('overall_mean_pairs', 'd_overall_pairs'), ('task_mean_pairs', 'd_task_mean_pairs'), ('allocation_mean_pairs', 'd_allocation_mean_pairs')]:
        shutil.copyfile(args.results_dir/f'{result_name}.tsv', args.output_dir/f'{source_name}.tsv')
    overall = read_tsv(args.output_dir/'d_overall_pairs.tsv')
    assert len(overall) == 18 and not overall[['depth_per_gemgroup', 'state_model']].duplicated().any()
    assert np.isfinite(overall[['d_S', 'V', 'd_D']]).all().all() and (overall.V > 0).all()
    assert set(overall.state_model) == set(STATE_MODELS)
    d = overall.copy()
    d['d_S_over_V'] = d.d_S / d.V
    d['inside_reversal_interval'] = (d.d_S_over_V > -1) & (d.d_S_over_V < 0)
    assert np.array_equal(d.inside_reversal_interval, d.expected_crossing)
    assert np.array_equal(d.expected_crossing, d.observed_crossing)
    assert np.max(np.abs(d.d_D - (d.d_S + d.V))) < 1e-12
    d.to_csv(args.output_dir/'d_summary.tsv', sep='\t', index=False, float_format='%.17g')
    tables = {path.stem: read_tsv(path) for path in sorted(args.output_dir.glob('*.tsv'))}
    assert len(tables['a_summary']) == 150 and len(tables['a_allocations']) == 4500
    assert len(tables['d_task_mean_pairs']) == 990 and len(tables['d_allocation_mean_pairs']) == 540
    notes = [
        ('Figure', 'Extended Data Figure 8: reference design in held-out genetic-combination prediction'),
        ('Scope', 'One pooled screen with eight technical capture groups. Tasks and allocations are not independent biological replicates.'),
        ('Training/test', '105 singles, 55 training doubles and ctrl; all 55 held-out doubles evaluated on the frozen 2,000-gene panel and scales.'),
        ('Effect model', 'Zero-intercept multihot ridge predicts native effects. Alpha 0.1 was selected using five training-condition folds before any held-out expression or evaluation controls were read.'),
        ('Models', 'Additive, compositional ridge and CPA are native-state predictors; direct-effect ridge predicts a fixed effect vector. Zero effect remains a diagnostic baseline.'),
        ('Controls', 'All 30 fixed allocations and six depths: 64, 128, 256, 512, 1024 and 1080 cells per role, divided equally among eight technical capture groups.'),
        ('a', 'Negative utility: average 55 task scores within allocation, then report allocation mean and full range. S and D are drawn; all five role patterns are retained.'),
        ('b', 'Unchanged CPA conditioning displacement: square root after averaging squared prediction differences across tasks, genes and ordered input-block pairs; then allocation mean/range.'),
        ('c', 'Unchanged original state-model comparisons: count strict crossings among 165 task/pair comparisons within each allocation, then summarize the 30 counts. Only P and O are drawn; S, M and D remain in source.'),
        ('d', 'All 18 effect/state/depth comparisons. Plot d_S/V after averaging tasks and allocations. Strict interval −1 < d_S/V < 0 is equivalent to −V < d_S < 0. Raw d_S, V and d_D remain in d_overall_pairs.'),
        ('d scale', 'First two facets share a y range; CPA uses a wider y range. No uncertainty interval is drawn for the completely enumerated aggregate comparisons.'),
        ('Pair orientation', 'Effect utility minus state utility; positive values favor the direct-effect predictor.'),
        ('Identity', 'd_D = d_S + V. All pair outcomes, including stable results, are retained.'),
        ('Variability', 'Bars in a–c are finite full allocation ranges, not confidence intervals.'),
        ('Aggregation order', 'Panel c counts before averaging allocations. crossings_after_mean instead averages task scores first. Panel d averages the paired margins before applying the reversal condition.'),
        ('Roles', 'S: all roles share; M: observation and prediction share; P: observation and model input share; O: prediction and model input share; D: all three roles distinct.'),
        ('Numerical ties', 'Margins with absolute value at most 1e-12 are treated as ties by the scorer. d_summary additionally retains the raw strict interval membership.'),
    ]
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = 'Reference-cell benchmark analysis'
    book.properties.title = 'Extended Data Figure 8 source data'
    book.properties.created = book.properties.modified = datetime(2000, 1, 1)
    for name, frame in {'Read me': pd.DataFrame(notes, columns=['Item', 'Description']), **tables}.items():
        sheet = book.create_sheet(name)
        sheet.append(frame.columns.tolist())
        for row in frame.itertuples(index=False, name=None):
            sheet.append([v.item() if isinstance(v, np.generic) else v for v in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='3D7180')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(14, len(str(column[0].value))+2))
    stable_xlsx(book, args.output_dir/'Extended_Data_Figure_8_Source_Data.xlsx')
    preserved = {name: sha256(args.output_dir/name) for name in source_names if not name.startswith('a_')}
    assert all(digest == original_record['output_sha256'][name] for name, digest in preserved.items())
    receipt = dict(status='PASS', biological_experiments=1, state_utility_shape=[55, 30, 6, 4, 5],
                   effect_utility_shape=list(utility.shape), original_state_rows_preserved_exactly=True,
                   unchanged_table_sha256=preserved, original_source_receipt_sha256=sha256(args.original_source_dir/'SOURCE_DATA_RECEIPT.json'),
                   source_utilities_sha256=original_record['source_utilities_sha256'], source_record_sha256=original_record['source_record_sha256'],
                   effect_score_receipt_sha256=sha256(args.results_dir/'score_receipt.json'),
                   effect_result_files_sha256=score_record['output_sha256'], protocol_sha256=score_record['protocol_sha256'],
                   tables={name: len(frame) for name, frame in tables.items()},
                   panel_d_pairs=18, panel_d_crossings=int(d.inside_reversal_interval.sum()),
                   panel_d_normalization='ratio of aggregate d_S and aggregate V; not an average of allocation-level ratios',
                   script_sha256=sha256(__file__),
                   output_sha256={p.name: sha256(p) for p in sorted(args.output_dir.iterdir()) if p.suffix in ('.tsv', '.xlsx')})
    (args.output_dir/'SOURCE_DATA_RECEIPT.json').write_text(json.dumps(receipt, indent=2)+'\n')
    print(json.dumps(receipt, indent=2))


if __name__ == '__main__':
    main()
